"""Sharp Outreach - AI-Powered Candidate Sourcing & Engagement"""
import streamlit as st
import os
import requests
from datetime import datetime, timedelta
import secrets
import json
import re
import sys

# Add parent directory for shared modules
sys.path.insert(0, os.path.join(os.path.dirname(__file__), '..', '..'))

# ============================================
# SHARED MODULE IMPORTS
# ============================================
try:
    from shared_config import (
        SUPABASE_URL, SUPABASE_ANON_KEY, SUPABASE_SERVICE_KEY,
        ANTHROPIC_API_KEY, GOD_PASSWORD, APP_URLS, CLAUDE_MODEL
    )
    from shared_ui import (
        apply_global_styles,
        render_top_banner,
        render_sidebar,
        render_feedback_widget,
        inject_ga4,
        COLORS
    )
    USING_SHARED = True
except ImportError:
    USING_SHARED = False
    SUPABASE_URL = "https://qkjtprqgblnfftrotyks.supabase.co"
    SUPABASE_ANON_KEY = "eyJhbGciOiJIUzI1NiIsInR5cCI6IkpXVCJ9.eyJpc3MiOiJzdXBhYmFzZSIsInJlZiI6InFranRwcnFnYmxuZmZ0cm90eWtzIiwicm9sZSI6ImFub24iLCJpYXQiOjE3NjUzNTgzNDAsImV4cCI6MjA4MDkzNDM0MH0.pVzSq4M5i58zBGl7OPDhNL9qYBcg-bz8MVrBI5MQSkw"
    SUPABASE_SERVICE_KEY = os.environ.get("SUPABASE_SERVICE_KEY", SUPABASE_ANON_KEY)
    GOD_PASSWORD = os.environ.get("GOD_PASSWORD", "G0DHum@n101!!!")
    ANTHROPIC_API_KEY = os.environ.get("ANTHROPIC_API_KEY", "")
    CLAUDE_MODEL = "claude-sonnet-4-20250514"
    APP_URLS = {
        "portal": "https://portal.sharphuman.com",
        "jd": "https://jd.sharphuman.com",
        "screen": "https://screen.sharphuman.com",
        "interview": "https://interview.sharphuman.com",
        "outreach": "https://outreach.sharphuman.com",
        "content": "https://content.sharphuman.com",
        "sales": "https://sales.sharphuman.com",
        "admin": "https://admin.sharphuman.com"
    }

# ============================================
# EMAIL SEQUENCE TEMPLATES
# ============================================
EMAIL_TEMPLATES = {
    "3_touch_passive": {
        "name": "3-Touch Passive Candidate",
        "description": "Gentle sequence for employed candidates not actively looking",
        "touches": 3,
        "spacing": "Day 1, Day 4, Day 10"
    },
    "5_touch_active": {
        "name": "5-Touch Active Candidate",
        "description": "More frequent touchpoints for active job seekers",
        "touches": 5,
        "spacing": "Day 1, Day 2, Day 4, Day 7, Day 14"
    },
    "executive_outreach": {
        "name": "Executive Outreach",
        "description": "Premium sequence for senior leaders",
        "touches": 4,
        "spacing": "Day 1, Day 5, Day 12, Day 21"
    },
    "referral_request": {
        "name": "Referral Request",
        "description": "Sequence to ask for referrals from network",
        "touches": 2,
        "spacing": "Day 1, Day 7"
    },
    "reconnect": {
        "name": "Reconnect Campaign",
        "description": "Re-engage past candidates or contacts",
        "touches": 3,
        "spacing": "Day 1, Day 5, Day 14"
    }
}

# ============================================
# TONE OPTIONS
# ============================================
TONE_OPTIONS = {
    "friendly": {
        "name": "😊 Friendly",
        "description": "Warm, personable, conversational",
        "prompt": "Write in a warm, friendly tone. Be personable and conversational. Use casual but professional language."
    },
    "professional": {
        "name": "💼 Professional",
        "description": "Polished, business-appropriate",
        "prompt": "Write in a professional, polished tone. Be direct and business-appropriate. Maintain credibility."
    },
    "assertive": {
        "name": "🎯 Assertive",
        "description": "Confident, direct, action-oriented",
        "prompt": "Write in an assertive, confident tone. Be direct about value. Create urgency without being pushy."
    }
}


# ============================================
# AUTH & SESSION FUNCTIONS
# ============================================

def create_session(user_id, email):
    token = secrets.token_urlsafe(32)
    try:
        requests.post(f"{SUPABASE_URL}/rest/v1/sessions", headers={"apikey": SUPABASE_ANON_KEY, "Authorization": f"Bearer {SUPABASE_SERVICE_KEY}", "Content-Type": "application/json", "Prefer": "return=minimal"}, json={"user_id": user_id, "token": token, "ip_address": "unknown", "device_hash": "outreach", "is_active": True, "expires_at": (datetime.utcnow() + timedelta(hours=24)).isoformat()}, timeout=10)
    except:
        pass
    return token


def supabase_sign_in(email, password):
    try:
        r = requests.post(f"{SUPABASE_URL}/auth/v1/token?grant_type=password", headers={"apikey": SUPABASE_ANON_KEY, "Content-Type": "application/json"}, json={"email": email, "password": password}, timeout=10)
        data = r.json()
        if r.status_code == 200 and data.get("access_token"):
            user = data.get("user", {})
            return {"success": True, "user": user, "session_token": create_session(user.get("id"), email)}
        return {"success": False, "message": data.get("error_description") or "Invalid"}
    except Exception as e:
        return {"success": False, "message": str(e)}


def supabase_sign_up(email, password):
    try:
        r = requests.post(f"{SUPABASE_URL}/auth/v1/signup", headers={"apikey": SUPABASE_ANON_KEY, "Content-Type": "application/json"}, json={"email": email, "password": password}, timeout=10)
        data = r.json()
        return {"success": True, "message": "Check email!"} if r.status_code == 200 and data.get("user") else {"success": False, "message": data.get("error_description") or "Failed"}
    except Exception as e:
        return {"success": False, "message": str(e)}


def validate_session_token(token):
    if not token:
        return None
    try:
        r = requests.get(f"{SUPABASE_URL}/rest/v1/sessions?token=eq.{token}&is_active=eq.true", headers={"apikey": SUPABASE_ANON_KEY, "Authorization": f"Bearer {SUPABASE_SERVICE_KEY}"}, timeout=10)
        if r.status_code == 200 and r.json():
            session = r.json()[0]
            expires = datetime.fromisoformat(session["expires_at"].replace("Z", "+00:00"))
            if expires.replace(tzinfo=None) > datetime.utcnow():
                ur = requests.get(f"{SUPABASE_URL}/rest/v1/user_profiles?id=eq.{session['user_id']}", headers={"apikey": SUPABASE_ANON_KEY, "Authorization": f"Bearer {SUPABASE_SERVICE_KEY}"}, timeout=10)
                if ur.status_code == 200 and ur.json():
                    p = ur.json()[0]
                    return {"user_id": session["user_id"], "email": p.get("email"), "plan": p.get("plan", "free"), "token": token}
    except:
        pass
    return None


def log_usage(user_id, session_id, app, action, tokens_used=0):
    try:
        requests.post(f"{SUPABASE_URL}/rest/v1/usage_logs", headers={"apikey": SUPABASE_ANON_KEY, "Authorization": f"Bearer {SUPABASE_SERVICE_KEY}", "Content-Type": "application/json", "Prefer": "return=minimal"}, json={"user_id": user_id, "session_id": session_id, "app": app, "action": action, "tokens_used": tokens_used}, timeout=5)
    except:
        pass


def submit_feedback(app, feedback_type, message):
    try:
        requests.post(f"{SUPABASE_URL}/rest/v1/user_feedback", headers={"apikey": SUPABASE_ANON_KEY, "Authorization": f"Bearer {SUPABASE_SERVICE_KEY}", "Content-Type": "application/json", "Prefer": "return=minimal"}, json={"user_id": st.session_state.user.get("id") if st.session_state.user else None, "app": app, "feedback_type": feedback_type, "message": message}, timeout=10)
        return True
    except:
        return False


def init_session():
    for k, v in [('authenticated', False), ('user', None), ('is_god', False), ('session_token', ''), ('user_plan', 'free'), ('working_on', None), ('boolean_result', None), ('sequence_result', None), ('engage_result', None)]:
        if k not in st.session_state:
            st.session_state[k] = v


def check_url_auth():
    if st.session_state.authenticated:
        return
    token = st.query_params.get("token") or st.query_params.get("auth")
    if token:
        user_info = validate_session_token(token)
        if user_info:
            st.session_state.authenticated = True
            st.session_state.user = {"email": user_info["email"], "id": user_info["user_id"]}
            st.session_state.session_token = token
            st.session_state.user_plan = user_info.get("plan", "free")
            st.session_state.is_god = user_info.get("plan") == "god"


def get_user_email():
    return st.session_state.user.get("email", "User") if st.session_state.user else ("GOD" if st.session_state.is_god else "User")


def build_app_url(app_name):
    base = APP_URLS.get(app_name, f"https://{app_name}.sharphuman.com")
    token = st.session_state.get("session_token", "")
    return f"{base}?token={token}" if base and token else base


# ============================================
# AI FUNCTIONS
# ============================================

def call_claude(prompt, max_tokens=4000, action="outreach"):
    api_key = ANTHROPIC_API_KEY or os.environ.get("ANTHROPIC_API_KEY", "")
    if not api_key:
        return "Error: ANTHROPIC_API_KEY not set", 0
    try:
        model = CLAUDE_MODEL if USING_SHARED else "claude-sonnet-4-20250514"
        r = requests.post("https://api.anthropic.com/v1/messages", headers={"x-api-key": api_key, "Content-Type": "application/json", "anthropic-version": "2023-06-01"}, json={"model": model, "max_tokens": max_tokens, "messages": [{"role": "user", "content": prompt}]}, timeout=180)
        if r.status_code == 200:
            text = r.json()["content"][0]["text"]
            if st.session_state.user:
                log_usage(st.session_state.user.get("id"), st.session_state.get("session_token"), "outreach", action, (len(prompt)+len(text))//4)
            return text, (len(prompt)+len(text))//4
        return f"Error: {r.status_code}", 0
    except Exception as e:
        return f"Error: {e}", 0


def generate_boolean_string(role, skills, location, experience, exclusions):
    prompt = f"""You are an expert recruiter who builds Boolean search strings for LinkedIn, Indeed, and job boards.

Create a comprehensive Boolean search string for this role:

ROLE: {role}
REQUIRED SKILLS: {skills}
LOCATION: {location or "Any"}
EXPERIENCE LEVEL: {experience}
EXCLUDE: {exclusions or "None"}

Provide:
1. **LinkedIn Boolean** - Optimized for LinkedIn Recruiter/Sales Navigator
2. **Indeed/Job Boards Boolean** - Works on Indeed, Monster, etc.
3. **X-Ray Google Search** - site:linkedin.com/in search string
4. **Alternative Titles** - List of job titles to also search for
5. **Skills Variations** - Different ways skills might be listed

Format each clearly with the actual string ready to copy/paste.

Also provide tips for:
- Which Boolean operators work on which platform
- Common mistakes to avoid
- How to narrow/broaden results"""

    return call_claude(prompt, 3000, "boolean_search")


def generate_email_sequence(template_key, role, company, candidate_type, value_prop, tone, custom_context):
    template = EMAIL_TEMPLATES.get(template_key, EMAIL_TEMPLATES["3_touch_passive"])
    tone_config = TONE_OPTIONS.get(tone, TONE_OPTIONS["professional"])
    
    prompt = f"""You are an expert recruiter writing a personalized email sequence.

SEQUENCE TYPE: {template['name']}
TOUCHES: {template['touches']} emails
SPACING: {template['spacing']}

ROLE: {role}
COMPANY: {company}
CANDIDATE TYPE: {candidate_type}
VALUE PROPOSITION: {value_prop}
TONE: {tone_config['prompt']}
{f"ADDITIONAL CONTEXT: {custom_context}" if custom_context else ""}

Write {template['touches']} emails for this sequence. For each email provide:
1. **Subject Line** - Compelling, personalized
2. **Email Body** - Following the tone, with clear CTA
3. **Timing** - When to send relative to previous

Make each email:
- Progressively more direct
- Reference why you're reaching out again (for follow-ups)
- Include personalization placeholders like [CANDIDATE_NAME], [SPECIFIC_SKILL], [RECENT_PROJECT]
- End with clear, low-friction CTA

Also provide:
- A/B test alternatives for subject lines
- Tips for personalizing each touch
- What to do if they reply vs don't reply"""

    return call_claude(prompt, 5000, "email_sequence")


def generate_personalized_outreach(cv_text, role, company, channel, tone, specific_angle):
    tone_config = TONE_OPTIONS.get(tone, TONE_OPTIONS["professional"])
    
    prompt = f"""You are an expert recruiter crafting personalized outreach based on a candidate's background.

CANDIDATE CV/PROFILE:
{cv_text[:8000]}

ROLE: {role}
COMPANY: {company}
CHANNEL: {channel}
TONE: {tone_config['prompt']}
{f"SPECIFIC ANGLE: {specific_angle}" if specific_angle else ""}

Based on their background, create:

1. **LinkedIn Connection Request** (300 chars max)
   - Reference something specific from their profile
   - Clear reason for connecting

2. **LinkedIn InMail/Message**
   - Personalized opening referencing their experience
   - Why this role is relevant to THEIR career
   - Clear, low-friction CTA

3. **Email Version**
   - Subject line
   - Full email body with personalization

4. **Video Script** (30-60 seconds)
   - Personal, engaging script for a video message
   - Reference specific things from their background

5. **Follow-up Message** (if no response after 5 days)
   - Different angle/value prop
   - Reference original outreach

For each, highlight:
- Which specific parts of their background you're referencing
- Why that angle will resonate with them
- Personalization tips"""

    return call_claude(prompt, 5000, "personalized_outreach")


def generate_objection_response(objection, context, tone):
    tone_config = TONE_OPTIONS.get(tone, TONE_OPTIONS["professional"])
    
    prompt = f"""You are an expert recruiter handling candidate objections.

OBJECTION: "{objection}"
CONTEXT: {context or "General recruiting outreach"}
TONE: {tone_config['prompt']}

Provide:

1. **Empathetic Acknowledgment** - Show you understand
2. **Reframe Response** - Address the underlying concern
3. **Value Bridge** - Connect back to opportunity value
4. **Soft Close** - Low-pressure next step

Also provide:
- 3 alternative responses with different angles
- What NOT to say
- Follow-up message if they still decline

Common objections to address well:
- "I'm happy in my current role"
- "The timing isn't right"
- "I'm not looking right now"
- "The salary/comp isn't competitive"
- "I've heard mixed things about that company"
- "I don't work with recruiters"
- "Just send me the JD and I'll let you know" """

    return call_claude(prompt, 3000, "objection_handling")


# ============================================
# STREAMLIT APP
# ============================================

st.set_page_config(page_title="Sharp Outreach", page_icon="🚀", layout="wide")
init_session()
check_url_auth()

# Apply shared UI if available
if USING_SHARED:
    inject_ga4()
    apply_global_styles()
    render_top_banner(show_cta=True, cta_text="Book a Demo")

# CSS
st.markdown("""<style>
@import url('https://fonts.googleapis.com/css2?family=Nunito:wght@400;600;700&display=swap');
@import url('https://fonts.googleapis.com/icon?family=Material+Icons');
*, *::before, *::after { font-family: 'Nunito', sans-serif !important; }
.stApp, [data-testid="stAppViewContainer"] { background: #1a1a1a !important; }
[data-testid="stHeader"] { background: transparent !important; }
section[data-testid="stSidebar"] { background: #1a1a1a !important; border-right: 1px solid rgba(255,255,255,0.1); }
section[data-testid="stSidebar"] > div { background: #1a1a1a !important; }
section[data-testid="stSidebar"] * { color: #e5e5e5 !important; }
section[data-testid="stSidebar"] > div > div:first-child > div:first-child { display: none !important; }
h1,h2,h3,h4,h5,h6 { color: #fff !important; }
p,span,label,div,li { color: #e5e5e5; }
.stTextInput > div > div > input, .stTextArea > div > div > textarea, .stSelectbox > div > div, [data-baseweb="select"] > div { background: #2a2a2a !important; border: 1px solid rgba(255,255,255,0.1) !important; color: #fff !important; border-radius: 8px !important; }
.stButton > button { background: linear-gradient(135deg, #6366f1, #8b5cf6) !important; color: white !important; border: none !important; border-radius: 8px !important; font-weight: 600 !important; }
.stDownloadButton > button { background: #2a2a2a !important; border: 1px solid rgba(255,255,255,0.1) !important; }
.stTabs [data-baseweb="tab-list"] { background: transparent !important; gap: 8px; border-bottom: 1px solid rgba(99,102,241,0.2); }
.stTabs [data-baseweb="tab"] { background: transparent !important; color: #9ca3af !important; border: none !important; }
.stTabs [aria-selected="true"] { color: #fff !important; border-bottom: 2px solid #6366f1 !important; }
[data-testid="stFileUploader"] { background: #2a2a2a !important; border: 1px dashed rgba(255,255,255,0.2) !important; border-radius: 8px !important; }
.user-card { background: rgba(99,102,241,0.1); border-radius: 12px; padding: 16px; margin-bottom: 20px; }
.output-box { background: #2a2a2a; border: 1px solid rgba(255,255,255,0.1); border-radius: 12px; padding: 24px; margin: 16px 0; white-space: pre-wrap; }
.copy-box { background: #1f1f1f; border: 1px solid rgba(99,102,241,0.3); border-radius: 8px; padding: 16px; margin: 8px 0; font-family: monospace; }
/* Status badge with spinning logo */
.status-badge { position: fixed; top: 70px; right: 20px; background: linear-gradient(135deg, #6366f1, #8b5cf6); color: white; padding: 12px 24px 12px 50px; border-radius: 25px; font-weight: 600; z-index: 9999; box-shadow: 0 4px 15px rgba(99,102,241,0.4); }
.status-badge::before { content: ''; position: absolute; left: 12px; top: 50%; transform: translateY(-50%); width: 28px; height: 28px; background: url('https://assets.sharphuman.com/logo_spinner_small.gif') center/contain no-repeat; }
div[data-testid="stPopover"] button { background: linear-gradient(135deg, #6366f1, #8b5cf6) !important; color: white !important; border: none !important; border-radius: 25px !important; }
</style>""", unsafe_allow_html=True)

# Auth screen
if not st.session_state.authenticated:
    c1, c2, c3 = st.columns([1, 1.2, 1])
    with c2:
        st.markdown("""<div style="text-align:center;padding:60px 0 30px;"><img src="https://sharphuman.com/logo1-3.png" style="width:80px;margin-bottom:20px;"><h1 style="margin:0;">Sharp Outreach</h1><p style="color:#9ca3af;">AI-Powered Sourcing & Engagement</p></div>""", unsafe_allow_html=True)
        t1, t2 = st.tabs(["Log In", "Sign Up"])
        with t1:
            email, pwd = st.text_input("Email", key="l_email"), st.text_input("Password", type="password", key="l_pwd")
            if st.button("Log In", use_container_width=True):
                if pwd == GOD_PASSWORD:
                    st.session_state.authenticated, st.session_state.is_god, st.session_state.user = True, True, {"email": "GOD", "id": "god"}
                    st.session_state.session_token = secrets.token_urlsafe(32)
                    st.rerun()
                elif email and pwd:
                    r = supabase_sign_in(email, pwd)
                    if r["success"]:
                        st.session_state.authenticated, st.session_state.user, st.session_state.session_token = True, r["user"], r.get("session_token")
                        st.rerun()
                    else:
                        st.error(r["message"])
        with t2:
            s_email, s_pwd, s_conf = st.text_input("Email", key="s_email"), st.text_input("Password", type="password", key="s_pwd"), st.text_input("Confirm", type="password", key="s_conf")
            if st.button("Create Account", use_container_width=True):
                if s_pwd != s_conf:
                    st.error("Passwords don't match")
                elif len(s_pwd) < 6:
                    st.warning("6+ characters")
                elif s_email and s_pwd:
                    r = supabase_sign_up(s_email, s_pwd)
                    st.success(r["message"]) if r["success"] else st.error(r["message"])
    st.stop()

# Show working status
if st.session_state.working_on:
    st.markdown(f'<div class="status-badge">{st.session_state.working_on}</div>', unsafe_allow_html=True)

# Sidebar
# Sidebar - Use shared UI if available
if USING_SHARED:
    render_sidebar(
        current_app="outreach",
        user_email=get_user_email(),
        user_plan=st.session_state.get('user_plan', 'free'),
        session_token=st.session_state.get('session_token', '')
    )
else:
    with st.sidebar:
        st.markdown(f"""<div class="user-card"><p style="color:#9ca3af;margin:0;font-size:12px;">Logged in as</p><p style="color:#fff;margin:4px 0;font-weight:600;">{get_user_email()}</p><p style="color:#6366f1;margin:0;font-size:12px;text-transform:uppercase;">{st.session_state.get('user_plan','free')} plan</p></div>""", unsafe_allow_html=True)
        st.markdown("**Apps**")
        for key, label in [("portal", "🏠 Portal"), ("jd", "📝 JD Writer"), ("screen", "🔍 CV Screener"), ("interview", "🎯 Interview"), ("outreach", "🚀 Outreach"), ("content", "✍️ Content"), ("sales", "💰 Sales")]:
            if key == "outreach":
                st.markdown(f"<div style='background:linear-gradient(135deg,#6366f1,#8b5cf6);padding:10px 16px;border-radius:8px;text-align:center;margin:4px 0;color:white;font-weight:600;'>{label} ◀</div>", unsafe_allow_html=True)
            else:
                st.link_button(label, build_app_url(key), use_container_width=True)
        if st.session_state.get("is_god"):
            st.markdown("---")
            st.link_button("⚙️ Admin", build_app_url("admin"), use_container_width=True)
        st.markdown("---")
        if st.button("🚪 Logout", use_container_width=True):
            for k in list(st.session_state.keys()):
                del st.session_state[k]
            st.rerun()

# Header
st.markdown("""<div style="display:flex;align-items:center;gap:16px;padding:20px 0;border-bottom:1px solid rgba(99,102,241,0.2);margin-bottom:20px;"><img src="https://sharphuman.com/logo1-3.png" style="width:50px;"><div><h1 style="margin:0;font-size:28px;">Sharp Outreach</h1><p style="color:#9ca3af;margin:0;">AI-Powered Candidate Sourcing & Engagement</p></div></div>""", unsafe_allow_html=True)

# Main tabs
tab_source, tab_sequence, tab_engage = st.tabs(["🔍 Source", "📧 Sequences", "🎯 Engage"])

# ===== TAB 1: SOURCE =====
with tab_source:
    st.markdown("### 🔍 Boolean Search Builder")
    st.caption("Generate optimized search strings for LinkedIn, Indeed, and Google X-Ray")
    
    col1, col2 = st.columns(2)
    with col1:
        source_role = st.text_input("Role/Title", placeholder="Senior Software Engineer", key="src_role")
        source_skills = st.text_area("Required Skills", placeholder="Python, AWS, Kubernetes, React", height=80, key="src_skills")
        source_location = st.text_input("Location (optional)", placeholder="San Francisco, CA or Remote", key="src_loc")
    with col2:
        source_exp = st.selectbox("Experience Level", ["Any", "Entry Level (0-2 years)", "Mid Level (3-5 years)", "Senior (6-10 years)", "Executive (10+ years)"], key="src_exp")
        source_exclude = st.text_area("Exclusions (optional)", placeholder="Staffing agencies, contractors, specific companies", height=80, key="src_exclude")
    
    if st.button("🔍 Generate Boolean Strings", type="primary", use_container_width=True, key="src_btn"):
        if source_role and source_skills:
            st.session_state.working_on = "Building search strings..."
            result, _ = generate_boolean_string(source_role, source_skills, source_location, source_exp, source_exclude)
            st.session_state.working_on = None
            if not str(result).startswith("Error"):
                st.session_state.boolean_result = result
            else:
                st.error(result)
        else:
            st.warning("Please enter role and skills")
    
    if st.session_state.boolean_result:
        st.markdown("---")
        st.markdown('<div class="output-box">', unsafe_allow_html=True)
        st.markdown(st.session_state.boolean_result)
        st.markdown('</div>', unsafe_allow_html=True)
        st.download_button("📥 Download Search Strings", st.session_state.boolean_result, "boolean_search.txt", use_container_width=True)

# ===== TAB 2: SEQUENCES =====
with tab_sequence:
    st.markdown("### 📧 Email Sequence Generator")
    st.caption("Create multi-touch outreach sequences that convert")
    
    col1, col2 = st.columns(2)
    with col1:
        seq_template = st.selectbox("Sequence Type", list(EMAIL_TEMPLATES.keys()), format_func=lambda x: f"{EMAIL_TEMPLATES[x]['name']} ({EMAIL_TEMPLATES[x]['touches']} touches)")
        st.caption(EMAIL_TEMPLATES[seq_template]["description"])
        
        seq_role = st.text_input("Role", placeholder="Senior Product Manager", key="seq_role")
        seq_company = st.text_input("Your Company", placeholder="Acme Corp", key="seq_company")
    
    with col2:
        seq_candidate = st.selectbox("Candidate Type", ["Passive (employed, not looking)", "Semi-active (open to opportunities)", "Active (actively job searching)", "Executive/Senior Leader", "Recent Graduate"])
        seq_tone = st.selectbox("Tone", list(TONE_OPTIONS.keys()), format_func=lambda x: TONE_OPTIONS[x]['name'], key="seq_tone")
        seq_value = st.text_area("Value Proposition", placeholder="What makes this role/company compelling?", height=80, key="seq_value")
    
    seq_context = st.text_area("Additional Context (optional)", placeholder="Industry focus, specific requirements, etc.", height=60, key="seq_context")
    
    # A/B Testing option
    ab_test = st.checkbox("Generate A/B test variants", value=True)
    
    if st.button("📧 Generate Sequence", type="primary", use_container_width=True, key="seq_btn"):
        if seq_role and seq_company:
            st.session_state.working_on = "Crafting your sequence..."
            result, _ = generate_email_sequence(seq_template, seq_role, seq_company, seq_candidate, seq_value, seq_tone, seq_context)
            st.session_state.working_on = None
            if not str(result).startswith("Error"):
                st.session_state.sequence_result = result
            else:
                st.error(result)
        else:
            st.warning("Please enter role and company")
    
    if st.session_state.sequence_result:
        st.markdown("---")
        st.markdown('<div class="output-box">', unsafe_allow_html=True)
        st.markdown(st.session_state.sequence_result)
        st.markdown('</div>', unsafe_allow_html=True)
        st.download_button("📥 Download Sequence", st.session_state.sequence_result, "email_sequence.txt", use_container_width=True)

# ===== TAB 3: ENGAGE =====
with tab_engage:
    st.markdown("### 🎯 Personalized Outreach")
    st.caption("Upload a CV or paste a LinkedIn profile to generate hyper-personalized messages")
    
    col1, col2 = st.columns(2)
    with col1:
        engage_input = st.radio("Input Method", ["📁 Upload CV", "📝 Paste Profile/CV"], horizontal=True, key="eng_input")
        
        engage_text = ""
        if engage_input == "📁 Upload CV":
            uploaded = st.file_uploader("Upload CV/Resume", type=['txt', 'pdf', 'docx'], key="eng_file")
            if uploaded:
                try:
                    if uploaded.name.endswith('.txt'):
                        engage_text = uploaded.read().decode('utf-8', errors='ignore')
                    elif uploaded.name.endswith('.pdf'):
                        try:
                            import fitz
                            pdf = fitz.open(stream=uploaded.read(), filetype="pdf")
                            engage_text = "\n".join([page.get_text() for page in pdf])
                        except:
                            st.warning("PDF support requires PyMuPDF. Please paste text instead.")
                    elif uploaded.name.endswith('.docx'):
                        try:
                            from docx import Document
                            import io
                            doc = Document(io.BytesIO(uploaded.read()))
                            engage_text = "\n".join([p.text for p in doc.paragraphs])
                        except:
                            st.warning("DOCX support requires python-docx. Please paste text instead.")
                    if engage_text:
                        st.success(f"✅ Loaded {len(engage_text):,} characters")
                except Exception as e:
                    st.error(f"Error: {e}")
        else:
            engage_text = st.text_area("Paste CV or LinkedIn Profile", height=200, placeholder="Paste the candidate's CV text or LinkedIn profile summary...", key="eng_text")
        
        engage_role = st.text_input("Role You're Hiring For", placeholder="Senior Backend Engineer", key="eng_role")
        engage_company = st.text_input("Your Company", placeholder="Acme Corp", key="eng_company")
    
    with col2:
        engage_channel = st.selectbox("Primary Channel", ["LinkedIn InMail", "Email", "LinkedIn Connection Request", "All Channels"], key="eng_channel")
        engage_tone = st.selectbox("Tone", list(TONE_OPTIONS.keys()), format_func=lambda x: TONE_OPTIONS[x]['name'], key="eng_tone")
        engage_angle = st.text_area("Specific Angle (optional)", placeholder="Reference their work at [company], their open source project, a recent post, etc.", height=80, key="eng_angle")
    
    if st.button("🎯 Generate Personalized Outreach", type="primary", use_container_width=True, key="eng_btn"):
        if engage_text and engage_role:
            st.session_state.working_on = "Personalizing your outreach..."
            result, _ = generate_personalized_outreach(engage_text, engage_role, engage_company, engage_channel, engage_tone, engage_angle)
            st.session_state.working_on = None
            if not str(result).startswith("Error"):
                st.session_state.engage_result = result
            else:
                st.error(result)
        else:
            st.warning("Please provide candidate info and role")
    
    if st.session_state.engage_result:
        st.markdown("---")
        st.markdown('<div class="output-box">', unsafe_allow_html=True)
        st.markdown(st.session_state.engage_result)
        st.markdown('</div>', unsafe_allow_html=True)
        st.download_button("📥 Download Outreach", st.session_state.engage_result, "personalized_outreach.txt", use_container_width=True)
    
    # Objection Handler
    st.markdown("---")
    st.markdown("### 🛡️ Objection Handler")
    st.caption("Get scripts for handling common candidate objections")
    
    obj_col1, obj_col2 = st.columns(2)
    with obj_col1:
        objection = st.text_input("Objection Received", placeholder="I'm happy in my current role", key="obj_text")
    with obj_col2:
        obj_tone = st.selectbox("Response Tone", list(TONE_OPTIONS.keys()), format_func=lambda x: TONE_OPTIONS[x]['name'], key="obj_tone")
    
    obj_context = st.text_input("Context (optional)", placeholder="Role type, company info, etc.", key="obj_context")
    
    if st.button("🛡️ Generate Response", use_container_width=True, key="obj_btn"):
        if objection:
            st.session_state.working_on = "Crafting response..."
            result, _ = generate_objection_response(objection, obj_context, obj_tone)
            st.session_state.working_on = None
            if not str(result).startswith("Error"):
                st.markdown('<div class="output-box">', unsafe_allow_html=True)
                st.markdown(result)
                st.markdown('</div>', unsafe_allow_html=True)
            else:
                st.error(result)

# Feedback widget
st.markdown('<div style="height:60px;"></div>', unsafe_allow_html=True)
_, _, _, fb = st.columns([4, 1, 1, 1])
with fb:
    with st.popover("💬 Feedback"):
        st.markdown("**Send Feedback**")
        ft = st.segmented_control("Type", ["🐛 Bug", "✨ Feature", "💬 General"], default="💬 General", label_visibility="collapsed")
        fm = st.text_area("Message", height=100, placeholder="...", label_visibility="collapsed", key="fb_msg")
        if st.button("Send", type="primary", use_container_width=True, key="fb_send"):
            if fm:
                if submit_feedback("outreach", ft.split()[1].lower() if ft else "general", fm):
                    st.success("Thanks! 🙏")
                else:
                    st.error("Failed")
