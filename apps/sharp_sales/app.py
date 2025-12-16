"""Sharp Sales - AI Sales Call Analysis (Enhanced)"""
import streamlit as st
import os
import requests
from datetime import datetime, timedelta, date
import secrets
import json
import re
import io
import tempfile

# Try to import few-shot examples
try:
    from sales_examples import get_sales_few_shot_examples
    HAS_EXAMPLES = True
except ImportError:
    HAS_EXAMPLES = False
    def get_sales_few_shot_examples():
        return ""

SUPABASE_URL = "https://qkjtprqgblnfftrotyks.supabase.co"
SUPABASE_ANON_KEY = "eyJhbGciOiJIUzI1NiIsInR5cCI6IkpXVCJ9.eyJpc3MiOiJzdXBhYmFzZSIsInJlZiI6InFranRwcnFnYmxuZmZ0cm90eWtzIiwicm9sZSI6ImFub24iLCJpYXQiOjE3NjUzNTgzNDAsImV4cCI6MjA4MDkzNDM0MH0.pVzSq4M5i58zBGl7OPDhNL9qYBcg-bz8MVrBI5MQSkw"
SUPABASE_SERVICE_KEY = os.environ.get("SUPABASE_SERVICE_KEY", SUPABASE_ANON_KEY)
GOD_PASSWORD = os.environ.get("GOD_PASSWORD", "G0DHum@n101!!!")
ANTHROPIC_API_KEY = os.environ.get("ANTHROPIC_API_KEY", "")

APP_URLS = {"portal": "https://demo.sharphuman.com", "jd": "https://jd.sharphuman.com", "screen": "https://screen.sharphuman.com", "interview": "https://interview.sharphuman.com", "source": "https://outreach.sharphuman.com", "content": "https://content.sharphuman.com", "sales": "https://sales.sharphuman.com", "reach": "https://reach.sharphuman.com", "assistant": "https://assistant.sharphuman.com", "admin": "https://admin.sharphuman.com"}

# ============================================
# COACH PERSONAS (Option 1)
# ============================================
COACH_PERSONAS = {
    "supportive": {
        "name": "🤝 Supportive Coach",
        "description": "Encouraging, focuses on wins first (good for new reps)",
        "system_prompt": "You are a supportive, encouraging sales coach. Lead with what went well before discussing improvements. Frame all feedback constructively. Your goal is to build confidence while guiding growth."
    },
    "balanced": {
        "name": "⚖️ Balanced Coach",
        "description": "Equal focus on strengths and improvements",
        "system_prompt": "You are a balanced sales coach providing fair, thorough analysis. Cover both strengths and areas for improvement equally. Be direct but constructive."
    },
    "savage": {
        "name": "🔥 Savage Sales Manager",
        "description": "No sugarcoating - brutal honesty (for experienced reps)",
        "system_prompt": "You are a savage Sales Manager with over 20 years experience. You don't sugarcoat - you tell it like it is. Every weakness is a revenue leak. Be brutally honest because mediocrity costs deals."
    },
    "deal_closer": {
        "name": "💰 Deal Closer",
        "description": "Focuses on what killed/saved the deal",
        "system_prompt": "You are a deal-obsessed sales coach. Your ONLY focus is: did this call move the deal forward or backward? Analyze every moment through the lens of 'did this help close or hurt the close?' Ignore nice-to-haves."
    },
    "vp_sales": {
        "name": "👔 VP of Sales",
        "description": "Strategic view - patterns, process, scalability",
        "system_prompt": "You are a VP of Sales reviewing this call. Focus on strategic patterns: Is this rep following process? Are they coachable? What systemic issues need addressing? Think about how this rep's patterns affect the whole team."
    }
}

# ============================================
# RIGOR LEVELS (Option 2)
# ============================================
RIGOR_LEVELS = {
    "coaching": {
        "name": "🌱 Coaching Mode",
        "description": "Focus on 2-3 key improvements (less overwhelming)",
        "prompt_modifier": "Focus on the TOP 2-3 most impactful improvements. Don't overwhelm with minor issues. This is for coaching, not comprehensive audit."
    },
    "balanced": {
        "name": "⚖️ Full Analysis",
        "description": "Complete analysis of all areas",
        "prompt_modifier": "Provide a complete, thorough analysis of all areas. Be comprehensive but prioritize the most important findings."
    },
    "brutal": {
        "name": "🔥 Brutal Audit",
        "description": "Find EVERY weakness (for top performers)",
        "prompt_modifier": "This is a BRUTAL audit. Find EVERY weakness, missed opportunity, and suboptimal moment. Nothing is too small to mention. The rep wants to hear it ALL so they can fix it. Be ruthless."
    }
}

# ============================================
# FEEDBACK REASONS (Option 5)
# ============================================
FEEDBACK_REASONS = {
    "thumbs_down": [
        ("missed_moment", "Missed key moment in the call"),
        ("wrong_score", "Score doesn't match performance"),
        ("bad_advice", "Coaching advice was off"),
        ("missed_objection", "Missed an objection"),
        ("generic", "Too generic, not specific enough"),
        ("hallucination", "Made up something not in transcript"),
        ("other", "Other")
    ],
    "thumbs_up": [
        ("accurate", "Spot-on analysis"),
        ("actionable", "Great actionable advice"),
        ("caught_issue", "Caught something I missed"),
        ("good_scripts", "Helpful fix scripts"),
        ("other", "Other")
    ]
}

CALL_FRAMEWORKS = {
    "discovery": {"name": "Discovery Call", "stages": {
        "opening": {"name": "Opening", "weight": 15, "skills": ["Rapport Building", "Setting the Frame", "Permission to Proceed"]},
        "discovery": {"name": "Discovery", "weight": 40, "skills": ["Uncovering Pain & Goals", "Probing Questions", "Cost of Inaction", "Timeline & Urgency", "Budget Discovery"]},
        "qualify": {"name": "Qualification", "weight": 20, "skills": ["Decision Authority", "Buying Process", "Competitive Landscape"]},
        "close": {"name": "Close", "weight": 25, "skills": ["Summary & Recap", "Clear Next Steps", "Getting Commitment", "Multi-threading"]}
    }},
    "demo": {"name": "Demo/Presentation", "stages": {
        "setup": {"name": "Setup", "weight": 15, "skills": ["Discovery Recap", "Demo Agenda", "Attendee Check"]},
        "demo": {"name": "Demonstration", "weight": 40, "skills": ["Tailored Demo", "Storytelling", "Audience Engagement", "Objection Handling"]},
        "value": {"name": "Value", "weight": 20, "skills": ["ROI Discussion", "Differentiation", "Social Proof"]},
        "close": {"name": "Close", "weight": 25, "skills": ["Temperature Check", "Surfacing Concerns", "Next Steps"]}
    }},
    "proposal": {"name": "Proposal Call", "stages": {
        "review": {"name": "Proposal Review", "weight": 30, "skills": ["Proposal Walkthrough", "Pricing Presentation", "Value Justification", "Addressing Questions"]},
        "concerns": {"name": "Concerns", "weight": 35, "skills": ["Objection Handling", "Risk Mitigation", "Competitor Comparison", "Stakeholder Concerns"]},
        "close": {"name": "Close", "weight": 35, "skills": ["Decision Timeline", "Next Steps", "Verbal Commitment", "Contract Process"]}
    }},
    "negotiation": {"name": "Negotiation Call", "stages": {
        "review": {"name": "Review", "weight": 20, "skills": ["Proposal Recap", "Value Reinforcement"]},
        "negotiate": {"name": "Negotiation", "weight": 50, "skills": ["Listening to Concerns", "Trading Value", "Anchoring", "Creative Solutions"]},
        "close": {"name": "Close", "weight": 30, "skills": ["Asking for Close", "Final Objections", "Paperwork Process"]}
    }},
    "interview": {"name": "Customer Interview", "stages": {
        "intro": {"name": "Introduction", "weight": 15, "skills": ["Rapport Building", "Setting Context", "Permission & Recording"]},
        "discovery": {"name": "Discovery", "weight": 50, "skills": ["Open-Ended Questions", "Active Listening", "Follow-up Probes", "Capturing Insights", "Pain Point Exploration"]},
        "close": {"name": "Wrap-Up", "weight": 35, "skills": ["Summary of Key Points", "Additional Questions", "Next Steps", "Thank You & Follow-up"]}
    }},
    "follow_up": {"name": "Follow-Up Call", "stages": {
        "reconnect": {"name": "Reconnect", "weight": 25, "skills": ["Context Reset", "Situation Changes", "Value Reminder"]},
        "advance": {"name": "Advance", "weight": 50, "skills": ["Address Objections", "New Information", "Creating Urgency"]},
        "commit": {"name": "Commitment", "weight": 25, "skills": ["Micro-Commitment", "Book Next Meeting", "Action Items"]}
    }}
}


def create_session(user_id, email):
    token = secrets.token_urlsafe(32)
    try: requests.post(f"{SUPABASE_URL}/rest/v1/sessions", headers={"apikey": SUPABASE_ANON_KEY, "Authorization": f"Bearer {SUPABASE_SERVICE_KEY}", "Content-Type": "application/json", "Prefer": "return=minimal"}, json={"user_id": user_id, "token": token, "ip_address": "unknown", "device_hash": "sales", "is_active": True, "expires_at": (datetime.utcnow() + timedelta(hours=24)).isoformat()}, timeout=10)
    except: pass
    return token

def supabase_sign_in(email, password):
    try:
        r = requests.post(f"{SUPABASE_URL}/auth/v1/token?grant_type=password", headers={"apikey": SUPABASE_ANON_KEY, "Content-Type": "application/json"}, json={"email": email, "password": password}, timeout=10)
        data = r.json()
        if r.status_code == 200 and data.get("access_token"):
            user = data.get("user", {})
            return {"success": True, "user": user, "session_token": create_session(user.get("id"), email)}
        return {"success": False, "message": data.get("error_description") or "Invalid"}
    except Exception as e: return {"success": False, "message": str(e)}

def supabase_sign_up(email, password):
    try:
        r = requests.post(f"{SUPABASE_URL}/auth/v1/signup", headers={"apikey": SUPABASE_ANON_KEY, "Content-Type": "application/json"}, json={"email": email, "password": password}, timeout=10)
        data = r.json()
        return {"success": True, "message": "Check email!"} if r.status_code == 200 and data.get("user") else {"success": False, "message": data.get("error_description") or "Failed"}
    except Exception as e: return {"success": False, "message": str(e)}

def validate_session_token(token):
    if not token: return None
    try:
        r = requests.get(f"{SUPABASE_URL}/rest/v1/sessions?token=eq.{token}&is_active=eq.true", headers={"apikey": SUPABASE_ANON_KEY, "Authorization": f"Bearer {SUPABASE_SERVICE_KEY}"}, timeout=10)
        if r.status_code == 200 and r.json():
            session = r.json()[0]
            expires = datetime.fromisoformat(session["expires_at"].replace("Z", "+00:00"))
            if expires.replace(tzinfo=None) > datetime.utcnow():
                ur = requests.get(f"{SUPABASE_URL}/rest/v1/user_profiles?id=eq.{session['user_id']}", headers={"apikey": SUPABASE_ANON_KEY, "Authorization": f"Bearer {SUPABASE_SERVICE_KEY}"}, timeout=10)
                if ur.status_code == 200 and ur.json():
                    p = ur.json()[0]
                    return {"user_id": session["user_id"], "email": p.get("email"), "plan": p.get("plan", "free"), "token": token}
    except: pass
    return None

def log_usage(user_id, session_id, app, action, tokens_used=0):
    try: requests.post(f"{SUPABASE_URL}/rest/v1/usage_logs", headers={"apikey": SUPABASE_ANON_KEY, "Authorization": f"Bearer {SUPABASE_SERVICE_KEY}", "Content-Type": "application/json", "Prefer": "return=minimal"}, json={"user_id": user_id, "session_id": session_id, "app": app, "action": action, "tokens_used": tokens_used}, timeout=5)
    except: pass

def submit_feedback(app, feedback_type, message):
    try:
        r = requests.post(f"{SUPABASE_URL}/rest/v1/user_feedback", headers={"apikey": SUPABASE_ANON_KEY, "Authorization": f"Bearer {SUPABASE_SERVICE_KEY}", "Content-Type": "application/json", "Prefer": "return=minimal"}, json={"user_id": st.session_state.user.get("id") if st.session_state.user else None, "app": app, "feedback_type": feedback_type, "rating": 4, "message": message, "email": get_user_email()}, timeout=10)
        return r.status_code in [200, 201]
    except: return False

def submit_analysis_feedback(result_data, feedback, reason=None, comment=None):
    """Submit thumbs up/down feedback for analysis quality"""
    try:
        payload = {
            "user_id": st.session_state.user.get("id") if st.session_state.user else None,
            "app": "sales",
            "persona": result_data.get("persona", "balanced"),
            "rigor_level": result_data.get("rigor", "balanced"),
            "overall_score": result_data.get("overall_score", 0),
            "feedback": feedback,
            "feedback_reason": reason,
            "feedback_comment": comment,
            "created_at": datetime.utcnow().isoformat()
        }
        r = requests.post(
            f"{SUPABASE_URL}/rest/v1/analysis_feedback",
            headers={
                "apikey": SUPABASE_ANON_KEY,
                "Authorization": f"Bearer {SUPABASE_SERVICE_KEY}",
                "Content-Type": "application/json",
                "Prefer": "return=minimal"
            },
            json=payload,
            timeout=10
        )
        return r.status_code in [200, 201]
    except:
        return False

def init_session():
    defaults = [
        ('authenticated', False), ('user', None), ('is_god', False), 
        ('session_token', None), ('user_plan', 'free'), ('working_on', None), 
        ('analysis_result', None), ('feedback_given', False),
        ('prospect_name', ''), ('company_name', '')
    ]
    for k, v in defaults:
        if k not in st.session_state: 
            st.session_state[k] = v

def check_url_auth():
    token = st.query_params.get("auth")
    if token and not st.session_state.authenticated:
        user_info = validate_session_token(token)
        if user_info:
            st.session_state.authenticated, st.session_state.user = True, {"email": user_info["email"], "id": user_info["user_id"]}
            st.session_state.session_token, st.session_state.user_plan = token, user_info.get("plan", "free")
            st.session_state.is_god = user_info.get("plan") == "god"
            return True
    return False

def get_user_email(): return st.session_state.user.get("email", "User") if st.session_state.user else ("GOD" if st.session_state.is_god else "User")
def build_app_url(app_name):
    base, token = APP_URLS.get(app_name, ""), st.session_state.get("session_token", "")
    return f"{base}?auth={token}" if base and token else base

def extract_text_from_file(uploaded_file):
    """Extract text from uploaded files with robust handling"""
    import zipfile
    from xml.etree import ElementTree
    
    def is_readable_text(text):
        if not text or len(text.strip()) < 10:
            return False
        alpha_chars = sum(1 for c in text if c.isalpha())
        return (alpha_chars / len(text) if len(text) > 0 else 0) > 0.3
    
    def clean_text(text):
        if not text:
            return ""
        cleaned = ''.join(c if c.isprintable() or c in '\n\r\t' else ' ' for c in text)
        cleaned = re.sub(r'[ \t]+', ' ', cleaned)
        return re.sub(r'\n{3,}', '\n\n', cleaned).strip()
    
    file_type = uploaded_file.name.split('.')[-1].lower()
    content = uploaded_file.read()
    uploaded_file.seek(0)
    
    if file_type == 'txt':
        return clean_text(content.decode('utf-8', errors='ignore'))
    
    elif file_type == 'pdf':
        extracted_text = ""
        
        # Try PyMuPDF with multiple methods
        try:
            import fitz
            pdf = fitz.open(stream=content, filetype="pdf")
            text_parts = []
            for page in pdf:
                page_text = page.get_text("text")
                if not page_text or not page_text.strip():
                    page_text = page.get_text("text", flags=fitz.TEXT_PRESERVE_WHITESPACE)
                if not page_text or not page_text.strip():
                    blocks = page.get_text("blocks")
                    page_text = "\n".join([b[4] for b in blocks if b[6] == 0])
                if page_text and page_text.strip():
                    text_parts.append(page_text)
            pdf.close()
            extracted_text = "\n".join(text_parts)
        except:
            pass
        
        if extracted_text and extracted_text.strip():
            cleaned = clean_text(extracted_text)
            if is_readable_text(cleaned):
                return cleaned
        
        # Try pdfplumber
        try:
            import pdfplumber
            with pdfplumber.open(io.BytesIO(content)) as pdf:
                text_parts = [page.extract_text() for page in pdf.pages if page.extract_text()]
                if text_parts:
                    cleaned = clean_text("\n".join(text_parts))
                    if is_readable_text(cleaned):
                        return cleaned
        except:
            pass
        
        # Last resort
        try:
            import fitz
            pdf = fitz.open(stream=content, filetype="pdf")
            text_parts = [page.get_text() for page in pdf]
            pdf.close()
            result = clean_text("\n".join(text_parts))
            if len(result) > 50:
                return result
        except:
            pass
        
        return "[PDF extraction failed - please paste the transcript directly]"
    
    elif file_type in ['docx', 'doc']:
        # Try python-docx first
        try:
            from docx import Document
            doc = Document(io.BytesIO(content))
            paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
            for table in doc.tables:
                for row in table.rows:
                    for cell in row.cells:
                        if cell.text.strip():
                            paragraphs.append(cell.text.strip())
            if paragraphs:
                return clean_text('\n\n'.join(paragraphs))
        except:
            pass
        
        # Fallback: raw XML from DOCX
        try:
            with zipfile.ZipFile(io.BytesIO(content)) as z:
                xml_content = z.read('word/document.xml')
                tree = ElementTree.fromstring(xml_content)
                texts = [elem.text for elem in tree.iter() if elem.text and elem.text.strip()]
                if texts:
                    return clean_text(' '.join(texts))
        except:
            pass
        
        return "[DOCX extraction failed - please paste the transcript directly]"
    
    elif file_type in ['vtt', 'srt']:
        text = content.decode('utf-8', errors='ignore')
        text = re.sub(r'\d{2}:\d{2}:\d{2}[.,]\d{3}\s*-->\s*\d{2}:\d{2}:\d{2}[.,]\d{3}', '', text)
        text = re.sub(r'WEBVTT.*?\n', '', text)
        text = re.sub(r'<[^>]+>', '', text)
        text = re.sub(r'^\d+$', '', text, flags=re.MULTILINE)
        lines = [line.strip() for line in text.split('\n') if line.strip()]
        return '\n'.join(lines)
    
    elif file_type in ['mp3', 'm4a', 'wav', 'mp4', 'webm', 'ogg']:
        openai_key = os.environ.get("OPENAI_API_KEY", "")
        if openai_key:
            try:
                with tempfile.NamedTemporaryFile(suffix=f".{file_type}", delete=False) as tmp:
                    tmp.write(content)
                    tmp_path = tmp.name
                with open(tmp_path, 'rb') as f:
                    r = requests.post("https://api.openai.com/v1/audio/transcriptions", headers={"Authorization": f"Bearer {openai_key}"}, files={"file": f}, data={"model": "whisper-1"}, timeout=300)
                os.unlink(tmp_path)
                if r.status_code == 200:
                    return r.json().get("text", "")
                return f"[Transcription error: {r.status_code}]"
            except Exception as e:
                return f"[Transcription error: {e}]"
        return "[Audio transcription requires OPENAI_API_KEY environment variable]"
    
    # Default: try to decode as text
    return content.decode('utf-8', errors='ignore')

def call_claude(prompt, max_tokens=8000, action="sales"):
    if not ANTHROPIC_API_KEY: return "Error: ANTHROPIC_API_KEY not set", 0
    try:
        r = requests.post("https://api.anthropic.com/v1/messages", headers={"x-api-key": ANTHROPIC_API_KEY, "Content-Type": "application/json", "anthropic-version": "2023-06-01"}, json={"model": "claude-sonnet-4-20250514", "max_tokens": max_tokens, "messages": [{"role": "user", "content": prompt}]}, timeout=180)
        if r.status_code == 200:
            text = r.json()["content"][0]["text"]
            if st.session_state.user: log_usage(st.session_state.user.get("id"), st.session_state.get("session_token"), "sales", action, (len(prompt)+len(text))//4)
            return text, (len(prompt)+len(text))//4
        return f"Error: {r.status_code}", 0
    except Exception as e: return f"Error: {e}", 0

def export_to_pdf(data, title="Sales Call Analysis"):
    """Export analysis to PDF using reportlab"""
    try:
        from reportlab.lib.pagesizes import letter
        from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
        from reportlab.lib.units import inch
        from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, HRFlowable
        from reportlab.lib.colors import HexColor
        from reportlab.lib.enums import TA_CENTER
        
        buffer = io.BytesIO()
        doc = SimpleDocTemplate(buffer, pagesize=letter, topMargin=0.5*inch, bottomMargin=0.5*inch)
        
        styles = getSampleStyleSheet()
        title_style = ParagraphStyle('Title', parent=styles['Title'], fontSize=22, textColor=HexColor('#6366f1'), spaceAfter=16)
        heading_style = ParagraphStyle('Heading', parent=styles['Heading1'], fontSize=14, textColor=HexColor('#374151'), spaceBefore=16, spaceAfter=8)
        subhead_style = ParagraphStyle('Subhead', parent=styles['Heading2'], fontSize=12, textColor=HexColor('#6366f1'), spaceBefore=10, spaceAfter=6)
        body_style = ParagraphStyle('Body', parent=styles['Normal'], fontSize=10, textColor=HexColor('#4b5563'), spaceAfter=6)
        bullet_style = ParagraphStyle('Bullet', parent=styles['Normal'], fontSize=10, textColor=HexColor('#4b5563'), leftIndent=20, spaceAfter=4)
        quote_style = ParagraphStyle('Quote', parent=styles['Normal'], fontSize=9, textColor=HexColor('#6366f1'), leftIndent=20, spaceAfter=4, fontName='Helvetica-Oblique')
        
        story = []
        
        # Title
        story.append(Paragraph(f"📊 {title}", title_style))
        story.append(HRFlowable(width="100%", thickness=1, color=HexColor('#e5e7eb')))
        story.append(Spacer(1, 12))
        
        if isinstance(data, dict):
            # Overall Score
            score = data.get('overall_score', 0)
            rec = data.get('overall_recommendation', 'N/A')
            score_color = '#10b981' if score >= 75 else '#eab308' if score >= 50 else '#ef4444'
            
            score_data = [['Overall Score', 'Recommendation'], [f"{score}/100", rec]]
            score_table = Table(score_data, colWidths=[2.5*inch, 4*inch])
            score_table.setStyle(TableStyle([
                ('BACKGROUND', (0, 0), (-1, 0), HexColor('#f3f4f6')),
                ('TEXTCOLOR', (0, 1), (0, 1), HexColor(score_color)),
                ('FONTSIZE', (0, 1), (-1, 1), 14),
                ('FONTNAME', (0, 0), (-1, -1), 'Helvetica-Bold'),
                ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
                ('GRID', (0, 0), (-1, -1), 0.5, HexColor('#e5e7eb')),
                ('TOPPADDING', (0, 0), (-1, -1), 10),
                ('BOTTOMPADDING', (0, 0), (-1, -1), 10),
            ]))
            story.append(score_table)
            story.append(Spacer(1, 16))
            
            # Key Wins
            story.append(Paragraph("Key Wins", heading_style))
            for w in data.get('key_wins', []):
                w_clean = ''.join(c if c.isprintable() else ' ' for c in str(w))
                story.append(Paragraph(f"✓ {w_clean}", bullet_style))
            
            # Critical Improvements
            story.append(Paragraph("Critical Improvements", heading_style))
            for i in data.get('critical_improvements', []):
                i_clean = ''.join(c if c.isprintable() else ' ' for c in str(i))
                story.append(Paragraph(f"• {i_clean}", bullet_style))
            
            # Stage Analysis
            story.append(Paragraph("Stage-by-Stage Analysis", heading_style))
            for stage in data.get('stage_analysis', []):
                stage_name = stage.get('stage_name', 'Unknown')
                stage_score = stage.get('stage_score', 0)
                story.append(Paragraph(f"{stage_name} — {stage_score}/10", subhead_style))
                
                for skill in stage.get('skills', []):
                    skill_name = skill.get('skill_name', '')
                    skill_score = skill.get('score', 0)
                    story.append(Paragraph(f"<b>{skill_name}:</b> {skill_score}/10", body_style))
                    for fb in skill.get('feedback', []):
                        fb_clean = ''.join(c if c.isprintable() else ' ' for c in str(fb))
                        story.append(Paragraph(f"→ {fb_clean}", bullet_style))
            
            # Deal Insights
            ins = data.get('deal_insights', {})
            if ins:
                story.append(Paragraph("Deal Insights", heading_style))
                story.append(Paragraph(f"<b>Deal Probability:</b> {ins.get('deal_probability', 'Unknown')}", body_style))
                
                if ins.get('buying_signals'):
                    story.append(Paragraph("<b>Buying Signals:</b>", body_style))
                    for s in ins['buying_signals']:
                        s_clean = ''.join(c if c.isprintable() else ' ' for c in str(s))
                        story.append(Paragraph(f"• {s_clean}", bullet_style))
                
                if ins.get('red_flags'):
                    story.append(Paragraph("<b>Red Flags:</b>", body_style))
                    for f in ins['red_flags']:
                        f_clean = ''.join(c if c.isprintable() else ' ' for c in str(f))
                        story.append(Paragraph(f"• {f_clean}", bullet_style))
            
            # Coaching Summary
            coaching = data.get('coaching_summary', '')
            if coaching:
                story.append(Paragraph("Coaching Summary", heading_style))
                coaching_clean = ''.join(c if c.isprintable() or c in '\n' else ' ' for c in str(coaching))
                story.append(Paragraph(coaching_clean, body_style))
        
        doc.build(story)
        buffer.seek(0)
        return buffer.getvalue()
    except Exception as e:
        return None

def analyze_call(call_type, prospect_name, company_name, deal_size, stage, notes, transcript, 
                 persona="balanced", rigor="balanced", key_concerns=""):
    """
    Analyze a sales call with customizable coaching settings.
    
    Args:
        persona: Key from COACH_PERSONAS
        rigor: Key from RIGOR_LEVELS
        key_concerns: Optional user-provided specific concerns to address
    """
    fw = CALL_FRAMEWORKS.get(call_type, CALL_FRAMEWORKS['discovery'])
    stages_desc = "\n".join([f"### {s['name']} (Weight: {s['weight']}%)\nSkills: {', '.join(s['skills'])}" for s in fw['stages'].values()])
    
    # Get persona and rigor settings
    persona_config = COACH_PERSONAS.get(persona, COACH_PERSONAS["balanced"])
    rigor_config = RIGOR_LEVELS.get(rigor, RIGOR_LEVELS["balanced"])
    
    # Build key concerns section
    concerns_section = ""
    if key_concerns and key_concerns.strip():
        concerns_section = f"""
## SPECIFIC CONCERNS TO ADDRESS
The user has specific concerns they want you to address in this analysis:
{key_concerns}

Make sure to explicitly address these concerns in your analysis with evidence from the transcript.
"""

    # Get few-shot examples
    few_shot_section = get_sales_few_shot_examples() if HAS_EXAMPLES else ""
    
    prompt = f"""{persona_config['system_prompt']}

{rigor_config['prompt_modifier']}

You are analyzing a {fw['name']}.

## CALL CONTEXT
- Prospect: {prospect_name or 'Unknown'} at {company_name or 'Unknown Company'}
- Deal Size: {deal_size or 'Unknown'}
- Stage: {stage or 'Unknown'}
- Notes: {notes or 'None'}

## CALL-TYPE-SPECIFIC EVALUATION (Stages)
{stages_desc}

## CORE SALES SKILLS FRAMEWORK (Universal - applies to ALL call types)
Rate the rep on these 6 universal sales skills that transfer across any call type:

1. **Call Control (0-10)**: Did the rep maintain structure and authority? Smooth transitions? Kept conversation on track?
2. **Discovery Depth (0-10)**: How well did they uncover the prospect's situation, goals, pain points, and gaps?
3. **Belief Shifting (0-10)**: Did they effectively reframe objections and position their offering as the logical solution?
4. **Objection Handling (0-10)**: Were objections pre-handled or post-handled effectively? Did they have good responses?
5. **Pitch Effectiveness (0-10)**: Clear value prop? Good demos/examples? Connected offering to prospect's specific goals?
6. **Closing Strength (0-10)**: Created urgency? Clear next steps? Asked for commitment? Used social proof?

{concerns_section}
{few_shot_section}

## TRANSCRIPT
{transcript[:20000]}

---

Analyze this call thoroughly. Provide TWO levels of analysis:
1. **Stage Analysis**: Specific to this {fw['name']} - what happened at each stage
2. **Core Sales Skills**: Universal skills that apply to ALL call types (for tracking improvement over time)

**CRITICAL REQUIREMENTS:**
1. Use EXACT timestamps from the transcript in [MM:SS] or [MM:SS-MM:SS] format
2. Include DIRECT QUOTES from the transcript
3. For each objection, specify: Pre-Handled? Post-Handled? Effectiveness Score
4. Include "Fix for Next Call" scripts for weak areas
5. Explain WHY something worked or didn't work

Return your analysis in this JSON format:
```json
{{
    "overall_score": <0-100>,
    "overall_summary": "<2-3 sentence executive summary>",
    "persona": "{persona}",
    "rigor": "{rigor}",
    "stages": [
        {{
            "stage_name": "<Stage Name>",
            "stage_score": <0-10>,
            "skills": [
                {{
                    "skill_name": "<Skill Name>",
                    "score": <0-10>,
                    "score_label": "<Excellent|Good|Needs Work|Missed>",
                    "what_worked": ["<specific thing that worked with timestamp and quote>"],
                    "what_needed_improvement": ["<specific thing to improve with timestamp>"],
                    "transcript_examples": ["<exact quote with [MM:SS] timestamp>"],
                    "fix_for_next_call": "<specific script or approach for next time>"
                }}
            ]
        }}
    ],
    "core_sales_skills": {{
        "call_control": {{
            "score": <0-10>,
            "summary": "<one sentence summary>",
            "evidence": ["<specific example with timestamp>"],
            "improvement": "<what to work on>"
        }},
        "discovery_depth": {{
            "score": <0-10>,
            "summary": "<one sentence summary>",
            "evidence": ["<specific example with timestamp>"],
            "improvement": "<what to work on>"
        }},
        "belief_shifting": {{
            "score": <0-10>,
            "summary": "<one sentence summary>",
            "evidence": ["<specific example with timestamp>"],
            "improvement": "<what to work on>"
        }},
        "objection_handling": {{
            "score": <0-10>,
            "summary": "<one sentence summary>",
            "evidence": ["<specific example with timestamp>"],
            "improvement": "<what to work on>"
        }},
        "pitch_effectiveness": {{
            "score": <0-10>,
            "summary": "<one sentence summary>",
            "evidence": ["<specific example with timestamp>"],
            "improvement": "<what to work on>"
        }},
        "closing_strength": {{
            "score": <0-10>,
            "summary": "<one sentence summary>",
            "evidence": ["<specific example with timestamp>"],
            "improvement": "<what to work on>"
        }},
        "total_score": <0-60>,
        "strongest_skill": "<which of the 6 skills was best>",
        "weakest_skill": "<which of the 6 skills needs most work>"
    }},
    "objections_breakdown": [
        {{
            "objection": "<what the prospect said>",
            "timestamp": "<[MM:SS]>",
            "pre_handled": <true/false>,
            "post_handled": <true/false>,
            "effectiveness_score": <0-10>,
            "what_worked": "<what the rep did well>",
            "what_needed_improvement": "<what could be better>",
            "fix_for_next_call": "<specific script to handle this objection better>"
        }}
    ],
    "key_wins": ["<something done well with timestamp>", "<another strength>"],
    "critical_improvements": ["<most important fix>", "<second priority>"],
    "deal_insights": {{
        "buying_signals": ["<signal identified with quote>"],
        "red_flags": ["<concern with evidence>"],
        "next_steps_suggested": ["<recommended action>"],
        "deal_probability": "<High|Medium|Low>",
        "deal_probability_reasoning": "<why>"
    }},
    "final_takeaways": {{
        "biggest_strength": "<the rep's best quality on this call>",
        "biggest_weakness": "<the most impactful thing to fix>",
        "game_changer_for_next_call": "<the ONE thing that would most improve results>"
    }},
    "coaching_summary": "<paragraph of personalized coaching advice>"
}}
```

Be specific. Use exact quotes with timestamps. Be constructive but honest."""

    return call_claude(prompt, 8000, f"analyze_{call_type}")

st.set_page_config(page_title="Sharp Sales", page_icon="💰", layout="wide")
init_session()
check_url_auth()

st.markdown("""<style>
@import url('https://fonts.googleapis.com/css2?family=Nunito:wght@400;600;700&display=swap');
*, *::before, *::after { font-family: 'Nunito', sans-serif !important; }
.stApp, [data-testid="stAppViewContainer"] { background: #0a0a0f !important; }
[data-testid="stHeader"] { background: transparent !important; }
section[data-testid="stSidebar"] { background: #0d0d14 !important; border-right: 1px solid rgba(99,102,241,0.2); }
section[data-testid="stSidebar"] > div { background: #0d0d14 !important; }
section[data-testid="stSidebar"] * { color: #e5e5e5 !important; }
section[data-testid="stSidebar"] > div > div:first-child > div:first-child { display: none !important; }
h1,h2,h3,h4,h5,h6 { color: #fff !important; }
p,span,label,div,li { color: #e5e5e5; }
.stTextInput > div > div > input, .stTextArea > div > div > textarea, .stSelectbox > div > div, [data-baseweb="select"] > div { background: #12121a !important; border: 1px solid rgba(99,102,241,0.3) !important; color: #fff !important; border-radius: 8px !important; }
.stButton > button { background: linear-gradient(135deg, #6366f1, #8b5cf6) !important; color: white !important; border: none !important; border-radius: 8px !important; font-weight: 600 !important; }
.stDownloadButton > button { background: #1a1a2e !important; border: 1px solid rgba(99,102,241,0.3) !important; }
.stRadio > div { flex-direction: row !important; gap: 8px; flex-wrap: wrap; }
.stRadio > div > label { background: #12121a !important; padding: 10px 16px !important; border-radius: 8px !important; border: 1px solid rgba(99,102,241,0.2) !important; }
[data-testid="stFileUploader"] { background: #12121a !important; border: 1px dashed rgba(99,102,241,0.3) !important; border-radius: 8px !important; }
.user-card { background: rgba(99,102,241,0.1); border-radius: 12px; padding: 16px; margin-bottom: 20px; }
.input-section { background: #12121a; border: 1px solid rgba(99,102,241,0.2); border-radius: 12px; padding: 20px; margin: 12px 0; }
.transcript-quote { background: #1a1a2e; border-left: 3px solid #6366f1; padding: 12px 16px; margin: 8px 0; border-radius: 0 8px 8px 0; font-style: italic; color: #a5b4fc; }
.improvement-box { background: rgba(16,185,129,0.1); border: 1px solid rgba(16,185,129,0.3); border-radius: 8px; padding: 16px; margin-top: 12px; }
.status-badge { position: fixed; top: 70px; right: 20px; background: linear-gradient(135deg, #6366f1, #8b5cf6); color: white; padding: 10px 20px; border-radius: 25px; font-weight: 600; z-index: 999; animation: pulse 2s infinite; }
@keyframes pulse { 0%, 100% { opacity: 1; } 50% { opacity: 0.7; } }
div[data-testid="stPopover"] button { background: linear-gradient(135deg, #6366f1, #8b5cf6) !important; color: white !important; border: none !important; border-radius: 25px !important; }
</style>""", unsafe_allow_html=True)

# Auth
if not st.session_state.authenticated:
    c1, c2, c3 = st.columns([1, 1.2, 1])
    with c2:
        st.markdown("""<div style="text-align:center;padding:60px 0 30px;"><img src="https://sharphuman.com/logo1-3.png" style="width:80px;margin-bottom:20px;"><h1 style="margin:0;">Sharp Sales</h1><p style="color:#9ca3af;">AI Sales Call Analysis</p></div>""", unsafe_allow_html=True)
        t1, t2 = st.tabs(["Log In", "Sign Up"])
        with t1:
            email, pwd = st.text_input("Email", key="l_email"), st.text_input("Password", type="password", key="l_pwd")
            if st.button("Log In", use_container_width=True):
                if pwd == GOD_PASSWORD:
                    st.session_state.authenticated, st.session_state.is_god, st.session_state.user = True, True, {"email": "GOD", "id": "god"}
                    st.session_state.session_token = secrets.token_urlsafe(32)
                    st.rerun()
                elif email and pwd:
                    r = supabase_sign_in(email, pwd)
                    if r["success"]: st.session_state.authenticated, st.session_state.user, st.session_state.session_token = True, r["user"], r.get("session_token"); st.rerun()
                    else: st.error(r["message"])
        with t2:
            s_email, s_pwd, s_conf = st.text_input("Email", key="s_email"), st.text_input("Password", type="password", key="s_pwd"), st.text_input("Confirm", type="password", key="s_conf")
            if st.button("Create Account", use_container_width=True):
                if s_pwd != s_conf: st.error("Passwords don't match")
                elif len(s_pwd) < 6: st.warning("6+ characters")
                elif s_email and s_pwd:
                    r = supabase_sign_up(s_email, s_pwd)
                    st.success(r["message"]) if r["success"] else st.error(r["message"])
    st.stop()

if st.session_state.working_on: st.markdown(f'<div class="status-badge">{st.session_state.working_on}</div>', unsafe_allow_html=True)

# Sidebar
with st.sidebar:
    st.markdown(f"""<div class="user-card"><p style="color:#9ca3af;margin:0;font-size:12px;">Logged in as</p><p style="color:#fff;margin:4px 0;font-weight:600;">{get_user_email()}</p><p style="color:#6366f1;margin:0;font-size:12px;text-transform:uppercase;">{st.session_state.get('user_plan','free')} plan</p></div>""", unsafe_allow_html=True)
    st.markdown("**Apps**")
    for key, label in [("portal", "🏠 Portal"), ("jd", "📝 JD Writer"), ("screen", "🔍 CV Screener"), ("interview", "🎯 Interview"), ("source", "🎣 Sourcing"), ("content", "✍️ Content"), ("sales", "💰 Sales"), ("reach", "🚀 Reach"), ("assistant", "🤖 Assistant")]:
        if key == "sales": st.markdown(f"<div style='background:linear-gradient(135deg,#6366f1,#8b5cf6);padding:10px 16px;border-radius:8px;text-align:center;margin:4px 0;color:white;font-weight:600;'>{label} ◀</div>", unsafe_allow_html=True)
        else: st.link_button(label, build_app_url(key), use_container_width=True)
    if st.session_state.get("is_god"): st.markdown("---"); st.link_button("⚙️ Admin", build_app_url("admin"), use_container_width=True)
    st.markdown("---")
    if st.button("🚪 Logout", use_container_width=True):
        for k in list(st.session_state.keys()): del st.session_state[k]
        st.rerun()

# Header
st.markdown("""<div style="display:flex;align-items:center;gap:16px;padding:20px 0;border-bottom:1px solid rgba(99,102,241,0.2);margin-bottom:30px;"><img src="https://sharphuman.com/logo1-3.png" style="width:50px;"><div><h1 style="margin:0;font-size:28px;">Sharp Sales</h1><p style="color:#9ca3af;margin:0;">AI-Powered Sales Call Analysis</p></div></div>""", unsafe_allow_html=True)

# Check for results
if st.session_state.get('analysis_result'):
    # Results View
    if st.button("← New Analysis", type="secondary"):
        st.session_state.analysis_result = None
        st.rerun()
    
    try:
        txt = st.session_state.analysis_result
        m = re.search(r'```json\s*(.*?)\s*```', txt, re.DOTALL)
        a = json.loads(m.group(1) if m else txt)
        
        score = a.get('overall_score', 0)
        color = "#10b981" if score >= 70 else "#eab308" if score >= 50 else "#ef4444"
        
        # Show persona/rigor badges
        persona_name = COACH_PERSONAS.get(a.get('persona', 'balanced'), {}).get('name', '')
        rigor_name = RIGOR_LEVELS.get(a.get('rigor', 'balanced'), {}).get('name', '')
        
        st.markdown(f"""<div style="background:linear-gradient(135deg,rgba(99,102,241,0.1),rgba(139,92,246,0.1));border-radius:16px;padding:30px;text-align:center;margin-bottom:24px;">
            <div style="margin-bottom:12px;">
                <span style="background:rgba(99,102,241,0.2);color:#a5b4fc;padding:4px 12px;border-radius:12px;font-size:12px;margin:4px;">{persona_name}</span>
                <span style="background:rgba(139,92,246,0.2);color:#c4b5fd;padding:4px 12px;border-radius:12px;font-size:12px;margin:4px;">{rigor_name}</span>
            </div>
            <p style="color:#9ca3af;margin:0;">OVERALL SCORE</p>
            <p style="color:{color};font-size:64px;font-weight:bold;margin:10px 0;">{score}<span style="font-size:24px;color:#6b7280;">/100</span></p>
            <p style="color:#e5e5e5;">{a.get('overall_summary', '')}</p>
        </div>""", unsafe_allow_html=True)
        
        c1, c2 = st.columns(2)
        with c1:
            st.markdown("#### ✅ Key Wins")
            for w in a.get('key_wins', []): st.markdown(f"<div style='background:rgba(16,185,129,0.1);border-left:3px solid #10b981;padding:12px;margin:8px 0;border-radius:0 8px 8px 0;'>{w}</div>", unsafe_allow_html=True)
        with c2:
            st.markdown("#### 🎯 Critical Improvements")
            for i in a.get('critical_improvements', []): st.markdown(f"<div style='background:rgba(239,68,68,0.1);border-left:3px solid #ef4444;padding:12px;margin:8px 0;border-radius:0 8px 8px 0;'>{i}</div>", unsafe_allow_html=True)
        
        st.markdown("---")
        st.markdown("## 📈 Stage Analysis")
        
        for stg in a.get('stages', []):
            with st.expander(f"**{stg.get('stage_name')}** — {stg.get('stage_score', 0)}/10"):
                for sk in stg.get('skills', []):
                    sc = sk.get('score', 0)
                    bc = "#10b981" if sc >= 8 else "#eab308" if sc >= 6 else "#f97316" if sc >= 4 else "#ef4444"
                    sym = "✓" if sc >= 6 else "−" if sc >= 4 else "✗"
                    
                    st.markdown(f"<div style='background:#12121a;border-radius:12px;padding:16px;margin:12px 0;border:1px solid rgba(99,102,241,0.2);'><div style='display:flex;justify-content:space-between;'><span style='font-weight:600;color:#fff;'>{sk.get('skill_name')}</span><span style='background:rgba(99,102,241,0.2);color:{bc};padding:6px 14px;border-radius:20px;font-size:13px;'>{sym} {sc}/10</span></div></div>", unsafe_allow_html=True)
                    
                    # What Worked
                    if sk.get('what_worked'):
                        st.markdown("**✅ What Worked:**")
                        for fb in sk.get('what_worked', []): st.markdown(f"→ {fb}")
                    
                    # What Needed Improvement
                    if sk.get('what_needed_improvement'):
                        st.markdown("**❌ What Needed Improvement:**")
                        for fb in sk.get('what_needed_improvement', []): st.markdown(f"→ {fb}")
                    
                    # Legacy feedback field
                    if sk.get('feedback'):
                        st.markdown("**Feedback:**")
                        for fb in sk.get('feedback', []): st.markdown(f"→ {fb}")
                    
                    if sk.get('transcript_examples'):
                        st.markdown("**📍 From the call:**")
                        for q in sk.get('transcript_examples', []): st.markdown(f'<div class="transcript-quote">"{q}"</div>', unsafe_allow_html=True)
                    
                    if sk.get('fix_for_next_call') or sk.get('improvement_example'):
                        fix = sk.get('fix_for_next_call') or sk.get('improvement_example')
                        st.markdown(f'<div class="improvement-box"><p style="color:#10b981;margin:0 0 8px;font-weight:600;">⚡ Fix for Next Call:</p><p style="color:#e5e5e5;margin:0;font-style:italic;">{fix}</p></div>', unsafe_allow_html=True)
        
        # Core Sales Skills Section (n8n Framework - NEW)
        core_skills = a.get('core_sales_skills', {})
        if core_skills:
            st.markdown("---")
            st.markdown("## 🎯 Core Sales Skills")
            st.caption("Universal skills that transfer across ALL call types — track these over time!")
            
            # Skills radar/summary
            total = core_skills.get('total_score', 0)
            total_color = "#10b981" if total >= 48 else "#eab308" if total >= 36 else "#ef4444"
            strongest = core_skills.get('strongest_skill', 'N/A')
            weakest = core_skills.get('weakest_skill', 'N/A')
            
            # Summary row
            sum_c1, sum_c2, sum_c3 = st.columns(3)
            with sum_c1:
                st.markdown(f"""
                <div style="background:#12121a;border-radius:12px;padding:20px;text-align:center;">
                    <p style="color:#9ca3af;margin:0;font-size:12px;">SKILLS TOTAL</p>
                    <p style="color:{total_color};font-size:36px;font-weight:bold;margin:8px 0;">{total}<span style="font-size:18px;color:#6b7280;">/60</span></p>
                </div>
                """, unsafe_allow_html=True)
            with sum_c2:
                st.markdown(f"""
                <div style="background:rgba(16,185,129,0.1);border:1px solid rgba(16,185,129,0.3);border-radius:12px;padding:16px;text-align:center;">
                    <p style="color:#10b981;margin:0;font-size:12px;">💪 STRONGEST</p>
                    <p style="color:#fff;font-size:16px;font-weight:600;margin:8px 0;">{strongest.replace('_', ' ').title()}</p>
                </div>
                """, unsafe_allow_html=True)
            with sum_c3:
                st.markdown(f"""
                <div style="background:rgba(239,68,68,0.1);border:1px solid rgba(239,68,68,0.3);border-radius:12px;padding:16px;text-align:center;">
                    <p style="color:#ef4444;margin:0;font-size:12px;">🎯 FOCUS AREA</p>
                    <p style="color:#fff;font-size:16px;font-weight:600;margin:8px 0;">{weakest.replace('_', ' ').title()}</p>
                </div>
                """, unsafe_allow_html=True)
            
            # Individual skill cards
            skill_labels = {
                'call_control': ('🎮', 'Call Control'),
                'discovery_depth': ('🔍', 'Discovery Depth'),
                'belief_shifting': ('🧠', 'Belief Shifting'),
                'objection_handling': ('🛡️', 'Objection Handling'),
                'pitch_effectiveness': ('🎯', 'Pitch Effectiveness'),
                'closing_strength': ('🏁', 'Closing Strength')
            }
            
            # Two rows of 3 skills each
            row1_skills = ['call_control', 'discovery_depth', 'belief_shifting']
            row2_skills = ['objection_handling', 'pitch_effectiveness', 'closing_strength']
            
            for skill_row in [row1_skills, row2_skills]:
                cols = st.columns(3)
                for idx, skill_key in enumerate(skill_row):
                    skill_data = core_skills.get(skill_key, {})
                    skill_score = skill_data.get('score', 0)
                    skill_color = "#10b981" if skill_score >= 8 else "#eab308" if skill_score >= 6 else "#f97316" if skill_score >= 4 else "#ef4444"
                    emoji, label = skill_labels.get(skill_key, ('📊', skill_key.replace('_', ' ').title()))
                    
                    with cols[idx]:
                        with st.expander(f"{emoji} **{label}** — {skill_score}/10"):
                            st.markdown(f"**Summary:** {skill_data.get('summary', 'N/A')}")
                            
                            evidence = skill_data.get('evidence', [])
                            if evidence:
                                st.markdown("**📍 Evidence:**")
                                for e in evidence:
                                    st.markdown(f'<div class="transcript-quote">{e}</div>', unsafe_allow_html=True)
                            
                            improvement = skill_data.get('improvement', '')
                            if improvement:
                                st.markdown(f'<div class="improvement-box"><p style="color:#10b981;margin:0 0 8px;font-weight:600;">💡 To Improve:</p><p style="color:#e5e5e5;margin:0;">{improvement}</p></div>', unsafe_allow_html=True)
        
        # Objections Breakdown (NEW)
        objections = a.get('objections_breakdown', [])
        if objections:
            st.markdown("---")
            st.markdown("## 🛠️ Objection Handling Breakdown")
            
            for i, obj in enumerate(objections, 1):
                pre = "✅" if obj.get('pre_handled') else "❌"
                post = "✅" if obj.get('post_handled') else "❌"
                eff_score = obj.get('effectiveness_score', 0)
                eff_color = "#10b981" if eff_score >= 8 else "#eab308" if eff_score >= 6 else "#ef4444"
                
                st.markdown(f"""
                <div style="background:#12121a;border:1px solid rgba(99,102,241,0.2);border-radius:12px;padding:20px;margin:12px 0;">
                    <div style="display:flex;justify-content:space-between;align-items:center;margin-bottom:12px;">
                        <span style="font-weight:600;color:#fff;">🚧 Objection {i}: {obj.get('objection', '')[:80]}...</span>
                        <span style="color:{eff_color};font-weight:600;">{eff_score}/10</span>
                    </div>
                    <p style="color:#9ca3af;font-size:13px;margin:4px 0;">📍 {obj.get('timestamp', 'N/A')}</p>
                    <div style="display:flex;gap:16px;margin:12px 0;">
                        <span style="color:#9ca3af;">Pre-Handled? {pre}</span>
                        <span style="color:#9ca3af;">Post-Handled? {post}</span>
                    </div>
                </div>
                """, unsafe_allow_html=True)
                
                if obj.get('what_worked'):
                    st.markdown(f"**✔️ What Worked:** {obj.get('what_worked')}")
                if obj.get('what_needed_improvement'):
                    st.markdown(f"**❌ Needed Improvement:** {obj.get('what_needed_improvement')}")
                if obj.get('fix_for_next_call'):
                    st.markdown(f'<div class="improvement-box"><p style="color:#10b981;margin:0 0 8px;font-weight:600;">⚡ Fix for Next Call:</p><p style="color:#e5e5e5;margin:0;font-style:italic;">"{obj.get("fix_for_next_call")}"</p></div>', unsafe_allow_html=True)
        
        st.markdown("---")
        st.markdown("## 💼 Deal Insights")
        
        ins = a.get('deal_insights', {})
        c1, c2, c3 = st.columns(3)
        with c1:
            prob = ins.get('deal_probability', 'Unknown')
            pc = "#10b981" if prob == "High" else "#eab308" if prob == "Medium" else "#ef4444"
            st.markdown(f"<div style='background:#12121a;border-radius:12px;padding:20px;text-align:center;'><p style='color:#9ca3af;margin:0;font-size:12px;'>DEAL PROBABILITY</p><p style='color:{pc};font-size:28px;font-weight:bold;margin:8px 0;'>{prob}</p><p style='color:#6b7280;font-size:11px;'>{ins.get('deal_probability_reasoning', '')[:80]}...</p></div>", unsafe_allow_html=True)
        with c2:
            st.markdown("**🟢 Buying Signals**")
            for s in ins.get('buying_signals', []): st.markdown(f"• {s}")
        with c3:
            st.markdown("**🔴 Red Flags**")
            for f in ins.get('red_flags', []): st.markdown(f"• {f}")
        
        st.markdown("#### 📋 Suggested Next Steps")
        for s in ins.get('next_steps_suggested', []): st.markdown(f"- [ ] {s}")
        
        # Final Takeaways (NEW)
        takeaways = a.get('final_takeaways', {})
        if takeaways:
            st.markdown("---")
            st.markdown("## 🚀 Final Takeaways")
            
            c1, c2, c3 = st.columns(3)
            with c1:
                st.markdown(f"""
                <div style="background:rgba(16,185,129,0.1);border:1px solid rgba(16,185,129,0.3);border-radius:12px;padding:16px;height:100%;">
                    <p style="color:#10b981;font-weight:600;margin:0 0 8px;">🔥 Biggest Strength</p>
                    <p style="color:#e5e5e5;margin:0;font-size:14px;">{takeaways.get('biggest_strength', 'N/A')}</p>
                </div>
                """, unsafe_allow_html=True)
            with c2:
                st.markdown(f"""
                <div style="background:rgba(239,68,68,0.1);border:1px solid rgba(239,68,68,0.3);border-radius:12px;padding:16px;height:100%;">
                    <p style="color:#ef4444;font-weight:600;margin:0 0 8px;">⚠️ Biggest Weakness</p>
                    <p style="color:#e5e5e5;margin:0;font-size:14px;">{takeaways.get('biggest_weakness', 'N/A')}</p>
                </div>
                """, unsafe_allow_html=True)
            with c3:
                st.markdown(f"""
                <div style="background:rgba(99,102,241,0.1);border:1px solid rgba(99,102,241,0.3);border-radius:12px;padding:16px;height:100%;">
                    <p style="color:#6366f1;font-weight:600;margin:0 0 8px;">🎯 Game-Changer</p>
                    <p style="color:#e5e5e5;margin:0;font-size:14px;">{takeaways.get('game_changer_for_next_call', 'N/A')}</p>
                </div>
                """, unsafe_allow_html=True)
        
        st.markdown("---")
        st.markdown("## 🎓 Coaching Summary")
        st.markdown(f"<div style='background:linear-gradient(135deg,rgba(99,102,241,0.1),rgba(139,92,246,0.1));border-radius:12px;padding:24px;border-left:4px solid #6366f1;'>{a.get('coaching_summary', '')}</div>", unsafe_allow_html=True)
        
        st.markdown("---")
        col1, col2, col3 = st.columns(3)
        with col1:
            st.download_button("📥 Markdown", txt, "sales_analysis.md", use_container_width=True)
        with col2:
            pdf_bytes = export_to_pdf(a, f"Sales Call: {st.session_state.get('prospect_name', 'Analysis')}")
            if pdf_bytes:
                st.download_button("📥 PDF", pdf_bytes, "sales_analysis.pdf", "application/pdf", use_container_width=True)
        with col3:
            st.download_button("📥 JSON", json.dumps(a, indent=2), "sales_analysis.json", use_container_width=True)
        
        # Feedback Section (NEW)
        st.markdown("---")
        st.markdown("### 💬 Rate This Analysis")
        
        if st.session_state.feedback_given:
            st.success("Thanks for your feedback! 🙏")
        else:
            st.caption("Your feedback helps improve our AI analysis")
            
            fb_col1, fb_col2, fb_col3 = st.columns([1, 1, 3])
            
            with fb_col1:
                if st.button("👍 Helpful", use_container_width=True, key="thumbs_up"):
                    if submit_analysis_feedback(a, "thumbs_up"):
                        st.session_state.feedback_given = True
                        st.rerun()
            
            with fb_col2:
                if st.button("👎 Not Helpful", use_container_width=True, key="thumbs_down"):
                    st.session_state.show_feedback_reason = True
            
            if st.session_state.get('show_feedback_reason'):
                reason_options = FEEDBACK_REASONS["thumbs_down"]
                selected_reason = st.radio(
                    "What went wrong?",
                    options=[r[0] for r in reason_options],
                    format_func=lambda x: dict(reason_options)[x],
                    label_visibility="collapsed"
                )
                comment = st.text_input("Details (optional):", placeholder="e.g., 'Missed the pricing objection at 12:30'")
                
                sub_col1, sub_col2 = st.columns(2)
                with sub_col1:
                    if st.button("Submit Feedback", type="primary", use_container_width=True):
                        if submit_analysis_feedback(a, "thumbs_down", selected_reason, comment):
                            st.session_state.feedback_given = True
                            st.session_state.show_feedback_reason = False
                            st.rerun()
                with sub_col2:
                    if st.button("Cancel", use_container_width=True):
                        st.session_state.show_feedback_reason = False
                        st.rerun()
        
        # New Analysis Button at bottom
        st.markdown("---")
        if st.button("🔄 Analyze Another Call", use_container_width=True):
            st.session_state.analysis_result = None
            st.session_state.feedback_given = False
            st.rerun()
        
    except Exception as e:
        st.error(f"Parse error: {e}")
        st.text(st.session_state.analysis_result)

else:
    # Input Form (Single Screen)
    st.markdown("### 📞 Call Details")
    
    col1, col2 = st.columns(2)
    
    with col1:
        st.markdown('<div class="input-section">', unsafe_allow_html=True)
        st.markdown("**Call Information**")
        call_type = st.selectbox("Call Type", list(CALL_FRAMEWORKS.keys()), format_func=lambda x: CALL_FRAMEWORKS[x]['name'])
        
        c1, c2 = st.columns(2)
        with c1:
            prospect_name = st.text_input("Prospect Name", placeholder="John Smith")
        with c2:
            company_name = st.text_input("Company", placeholder="Acme Corp")
        
        c1, c2 = st.columns(2)
        with c1:
            deal_size = st.text_input("Deal Size", placeholder="$50,000")
        with c2:
            stage = st.selectbox("Sales Stage", ["Discovery", "Qualification", "Demo", "Proposal", "Negotiation", "Closed"])
        
        notes = st.text_area("Notes/Questions (optional)", height=80, placeholder="Any specific areas you want analyzed?")
        st.markdown('</div>', unsafe_allow_html=True)
    
    with col2:
        st.markdown('<div class="input-section">', unsafe_allow_html=True)
        st.markdown("**Call Recording / Transcript**")
        
        input_method = st.radio("Input Method:", ["📝 Paste Transcript", "📁 Upload File"], horizontal=True)
        
        transcript = ""
        
        if input_method == "📝 Paste Transcript":
            transcript = st.text_area(
                "Paste your transcript here:",
                height=200,
                placeholder="Salesperson: Hi John, thanks for taking the time today...\nProspect: Of course, happy to chat...\n\nPaste the full conversation transcript here."
            )
        else:
            uploaded_file = st.file_uploader(
                "Upload recording or transcript",
                type=['txt', 'pdf', 'docx', 'doc', 'vtt', 'srt', 'mp3', 'wav', 'm4a', 'mp4', 'webm'],
                help="Supports: TXT, PDF, DOCX, VTT, SRT, MP3, WAV, M4A, MP4, WebM"
            )
            if uploaded_file:
                with st.spinner(f"Processing {uploaded_file.name}..."):
                    transcript = extract_text_from_file(uploaded_file)
                
                if transcript and not transcript.startswith("["):
                    st.success(f"✅ Loaded: {uploaded_file.name} ({len(transcript):,} characters)")
                    with st.expander("Preview transcript"):
                        st.text(transcript[:2000] + ("..." if len(transcript) > 2000 else ""))
                else:
                    st.warning(transcript)
                    st.info("💡 If upload fails, try pasting the transcript directly instead.")
        
        st.markdown('</div>', unsafe_allow_html=True)
        
        # Analysis Settings
        st.markdown('<div class="input-section">', unsafe_allow_html=True)
        st.markdown("**⚙️ Analysis Settings**")
        
        # Coach Persona
        persona_options = {v["name"]: k for k, v in COACH_PERSONAS.items()}
        selected_persona_name = st.selectbox(
            "Coach Style:",
            options=list(persona_options.keys()),
            index=1,  # Default to Balanced
            help="Different coaching styles for different needs"
        )
        selected_persona = persona_options[selected_persona_name]
        st.caption(COACH_PERSONAS[selected_persona]["description"])
        
        # Rigor Level
        rigor_options = {v["name"]: k for k, v in RIGOR_LEVELS.items()}
        selected_rigor_name = st.select_slider(
            "Analysis Depth:",
            options=list(rigor_options.keys()),
            value="⚖️ Full Analysis"
        )
        selected_rigor = rigor_options[selected_rigor_name]
        
        # Key Concerns
        key_concerns = st.text_area(
            "Specific concerns (optional):",
            height=68,
            placeholder="e.g., 'Did they handle the pricing objection well?' or 'Watch for signs of multi-threading'"
        )
        
        st.markdown('</div>', unsafe_allow_html=True)
    
    # Analyze Button
    st.markdown("---")
    
    if st.button("🚀 Analyze Call", type="primary", use_container_width=True):
        if not transcript or len(transcript.strip()) < 100:
            st.warning("Please provide a transcript (paste or upload). Minimum 100 characters required.")
        else:
            # Store names for later use
            st.session_state.prospect_name = prospect_name
            st.session_state.company_name = company_name
            st.session_state.feedback_given = False
            
            # Show immediate spinner
            with st.spinner(f"🔄 Analyzing call as {selected_persona_name}... This takes 20-40 seconds."):
                st.session_state.working_on = f"Analyzing as {selected_persona_name}..."
                result, _ = analyze_call(
                    call_type, prospect_name, company_name, deal_size, stage, notes, transcript,
                    persona=selected_persona,
                    rigor=selected_rigor,
                    key_concerns=key_concerns
                )
                st.session_state.working_on = None
            
            if not str(result).startswith("Error"):
                st.session_state.analysis_result = result
                st.rerun()
            else:
                st.error(result)

# Feedback
st.markdown('<div style="height:60px;"></div>', unsafe_allow_html=True)
_, _, _, fb = st.columns([4, 1, 1, 1])
with fb:
    with st.popover("💬 Feedback"):
        st.markdown("**Send Feedback**")
        ft = st.segmented_control("Type", ["🐛 Bug", "✨ Feature", "💬 General"], default="💬 General", label_visibility="collapsed")
        fm = st.text_area("Message", height=100, placeholder="...", label_visibility="collapsed", key="fb_msg")
        if st.button("Send", type="primary", use_container_width=True, key="fb_send"):
            if fm:
                if submit_feedback("sales", ft.split()[1].lower() if ft else "general", fm): st.success("Thanks! 🙏")
                else: st.error("Failed")
